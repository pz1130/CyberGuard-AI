"""Knowledge base service: chunking, embedding, and semantic search."""
import io
import json
import logging
import hashlib
from typing import List, Dict, Any, Optional

from sqlalchemy import select, text
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.knowledge import (
    KnowledgeBase, Document, DocumentChunk,
    EMBEDDING_COLUMN_BY_DIM, SUPPORTED_EMBEDDING_DIMS,
    EMBEDDING_DIM_LARGE,
)

# SQL type cast used in pgvector ANN queries. Must match the column type
# declared in the model (vector for 1536, halfvec for 3072).
_SQL_VECTOR_TYPE_BY_DIM: Dict[int, str] = {
    1536: "vector",
    EMBEDDING_DIM_LARGE: "halfvec",
}
from app.services.llm_router import get_llm_router

logger = logging.getLogger(__name__)


# --------------------------------------------------------------------------
# Binary → text extraction
# --------------------------------------------------------------------------
def extract_text(raw: bytes, mime_type: Optional[str], filename: str = "") -> str:
    """Extract plain text from PDF / docx / utf-8 byte payloads.

    Routing:
      - application/pdf                                  → pypdf
      - application/vnd.openxmlformats-...wordprocessingml → python-docx
      - everything else                                  → utf-8 decode

    Raises ValueError on parse failure so the router returns a 400.
    """
    mt = (mime_type or "").lower()
    fn = (filename or "").lower()

    if mt == "application/pdf" or fn.endswith(".pdf"):
        try:
            from pypdf import PdfReader
        except ImportError as e:
            raise ValueError(f"PDF support requires pypdf: {e}")
        try:
            reader = PdfReader(io.BytesIO(raw))
            pages = []
            for i, page in enumerate(reader.pages):
                try:
                    pages.append(page.extract_text() or "")
                except Exception as pe:
                    logger.warning("[extract_text] PDF page %d failed: %s", i, pe)
            text_out = "\n\n".join(p.strip() for p in pages if p and p.strip())
            if not text_out:
                raise ValueError("PDF contains no extractable text (likely scanned image).")
            return text_out
        except ValueError:
            raise
        except Exception as e:
            raise ValueError(f"PDF parse failed: {e}")

    if mt in (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/msword",
    ) or fn.endswith(".docx"):
        try:
            from docx import Document as DocxDocument
        except ImportError as e:
            raise ValueError(f"docx support requires python-docx: {e}")
        try:
            doc = DocxDocument(io.BytesIO(raw))
            paragraphs = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
            # Also pull cell text from tables.
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text and cell.text.strip():
                            paragraphs.append(cell.text.strip())
            text_out = "\n\n".join(paragraphs).strip()
            if not text_out:
                raise ValueError("docx contains no text.")
            return text_out
        except ValueError:
            raise
        except Exception as e:
            raise ValueError(f"docx parse failed: {e}")

    # Default: UTF-8 text (.txt, .md, .csv, .json, .html, etc.)
    try:
        return raw.decode("utf-8")
    except UnicodeDecodeError:
        raise ValueError(
            "Unsupported file: not PDF/docx and not valid UTF-8. "
            "Convert binary formats to text first."
        )


# --------------------------------------------------------------------------
# Chunking
# --------------------------------------------------------------------------
def chunk_text(text: str, chunk_size: int = 500, overlap: int = 50) -> List[str]:
    """
    Split text into overlapping chunks. Operates on characters for simplicity.

    Args:
        text: input text
        chunk_size: max characters per chunk
        overlap: characters of overlap between consecutive chunks

    Returns:
        list of chunk strings (non-empty)
    """
    text = (text or "").strip()
    if not text:
        return []
    if len(text) <= chunk_size:
        return [text]

    chunks: List[str] = []
    start = 0
    step = max(chunk_size - overlap, 1)
    while start < len(text):
        end = min(start + chunk_size, len(text))
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= len(text):
            break
        start += step
    return chunks


# --------------------------------------------------------------------------
# Service
# --------------------------------------------------------------------------
class KnowledgeService:
    """Ingestion and retrieval against the knowledge base.

    Uses pgvector HNSW index on document_chunks.embedding for ANN search.
    Cosine similarity = 1 - vector_cosine_distance.
    """

    def __init__(self):
        self.router = get_llm_router()

    async def _embed(
        self,
        texts: List[str],
        embedding_model: Optional[str],
        provider_id: Optional[int] = None,
    ) -> List[List[float]]:
        """Embed a list of texts via the configured AI Provider."""
        return await self.router.embed(
            texts=texts,
            model=embedding_model,
            provider_id=provider_id,
        )

    async def ingest_document(
        self,
        db: AsyncSession,
        kb_id: int,
        filename: str,
        content: str,
        mime_type: Optional[str] = None,
        provider_id: Optional[int] = None,
    ) -> Document:
        """
        Chunk + embed text and persist as Document + DocumentChunk rows.

        Embeddings are written to the column matching KB.embedding_dim
        (`embedding` for 1536, `embedding_large` for 3072); the other column
        is left NULL. The CK constraint on document_chunks enforces XOR.
        """
        kb = await db.get(KnowledgeBase, kb_id)
        if not kb:
            raise ValueError(f"Knowledge base {kb_id} not found")

        col_name = EMBEDDING_COLUMN_BY_DIM.get(kb.embedding_dim)
        if not col_name:
            raise ValueError(
                f"KB {kb_id} has unsupported embedding_dim={kb.embedding_dim}. "
                f"Supported: {SUPPORTED_EMBEDDING_DIMS}"
            )

        chunks = chunk_text(content)
        if not chunks:
            raise ValueError("Document is empty after chunking")

        embeddings = await self._embed(
            texts=chunks,
            embedding_model=kb.embedding_model,
            provider_id=provider_id,
        )

        if len(embeddings) != len(chunks):
            raise ValueError(
                f"Embedding count mismatch: expected {len(chunks)}, got {len(embeddings)}"
            )

        # Validate the provider returned the dimension the KB expects.
        first_dim = len(embeddings[0]) if embeddings else 0
        if first_dim != kb.embedding_dim:
            raise ValueError(
                f"Provider returned {first_dim}-dim vectors but KB {kb_id} expects "
                f"{kb.embedding_dim}-dim. Either reconfigure the AI Provider's embedding "
                f"model or create a new KB with embedding_dim={first_dim}."
            )

        file_hash = hashlib.sha256(content.encode("utf-8")).hexdigest()

        doc = Document(
            kb_id=kb_id,
            filename=filename,
            content_chunks_json=None,  # Deprecated — vectors now in document_chunks
            file_hash=file_hash,
            file_size=len(content.encode("utf-8")),
            mime_type=mime_type,
            metadata_json={"chunk_count": len(chunks), "embedding_dim": kb.embedding_dim},
        )
        db.add(doc)
        await db.flush()  # Get doc.id without committing

        # Write to the dim-appropriate column; leave the other NULL.
        chunk_records: List[DocumentChunk] = []
        for i, (chunk, emb) in enumerate(zip(chunks, embeddings)):
            kwargs: Dict[str, Any] = {
                "document_id": doc.id,
                "kb_id": kb_id,
                "chunk_index": i,
                "content": chunk,
                "metadata_json": {"filename": filename},
            }
            kwargs[col_name] = emb
            chunk_records.append(DocumentChunk(**kwargs))

        db.add_all(chunk_records)
        await db.commit()
        await db.refresh(doc)
        return doc

    async def query(
        self,
        db: AsyncSession,
        kb_id: int,
        query: str,
        top_k: int = 5,
        similarity_threshold: float = 0.0,
        provider_id: Optional[int] = None,
    ) -> List[Dict[str, Any]]:
        """
        Embed query and return top-k most similar chunks using pgvector HNSW ANN.

        Cosine distance is used (vector_cosine_ops in HNSW); similarity =
        1 - cosine_distance.  Results below similarity_threshold are filtered.

        Returns list of dicts with keys: document_id, filename, chunk_index,
        content, score.
        """
        kb = await db.get(KnowledgeBase, kb_id)
        if not kb:
            raise ValueError(f"Knowledge base {kb_id} not found")

        query_embeddings = await self._embed(
            texts=[query],
            embedding_model=kb.embedding_model,
            provider_id=provider_id,
        )
        if not query_embeddings:
            return []

        q_vec = query_embeddings[0]
        if len(q_vec) != kb.embedding_dim:
            raise ValueError(
                f"Query embedding is {len(q_vec)}-dim but KB {kb_id} expects "
                f"{kb.embedding_dim}-dim. Check the AI Provider's embedding model."
            )

        col_name = EMBEDDING_COLUMN_BY_DIM.get(kb.embedding_dim)
        cast_type = _SQL_VECTOR_TYPE_BY_DIM.get(kb.embedding_dim)
        if not col_name or not cast_type:
            raise ValueError(
                f"KB {kb_id} has unsupported embedding_dim={kb.embedding_dim}"
            )

        # Use raw SQL for the ANN query — pgvector requires an explicit
        # ::vector(<dim>) or ::halfvec(<dim>) cast matching the column type.
        # `col_name` and `cast_type` are allowlisted from server-side maps,
        # so f-string interpolation is safe (no user input).
        import json as _json
        q_vec_json = _json.dumps(q_vec)

        raw_sql = text(f"""
            SELECT
                dc.document_id,
                dc.chunk_index,
                dc.content,
                (1 - (dc.{col_name} <=> (:q_vec)::{cast_type}({kb.embedding_dim}))) AS score
            FROM document_chunks dc
            WHERE dc.kb_id = :kb_id AND dc.{col_name} IS NOT NULL
            ORDER BY dc.{col_name} <=> (:q_vec)::{cast_type}({kb.embedding_dim})
            LIMIT :top_k
        """)

        result = await db.execute(
            raw_sql,
            {"q_vec": q_vec_json, "kb_id": kb_id, "top_k": top_k},
        )
        rows = result.fetchall()

        if not rows:
            return []

        # Fetch filenames in a single query
        doc_ids = list({r.document_id for r in rows})
        docs_result = await db.execute(
            select(Document.id, Document.filename).where(Document.id.in_(doc_ids))
        )
        doc_names = {row[0]: row[1] for row in docs_result.all()}

        results = []
        for row in rows:
            score: float = row.score  # type: ignore[assignment]
            if score < similarity_threshold:
                continue
            results.append({
                "document_id": row.document_id,
                "filename": doc_names.get(row.document_id, "unknown"),
                "chunk_index": row.chunk_index,
                "content": row.content,
                "score": round(score, 4),
            })

        return results

    async def delete_document_chunks(self, db: AsyncSession, document_id: int) -> int:
        """Delete all chunks for a document. Called on document deletion."""
        result = await db.execute(
            select(DocumentChunk)
            .where(DocumentChunk.document_id == document_id)
        )
        chunks = result.scalars().all()
        count = len(chunks)
        for chunk in chunks:
            await db.delete(chunk)
        await db.commit()
        return count


_knowledge_service: Optional[KnowledgeService] = None


def get_knowledge_service() -> KnowledgeService:
    global _knowledge_service
    if _knowledge_service is None:
        _knowledge_service = KnowledgeService()
    return _knowledge_service
